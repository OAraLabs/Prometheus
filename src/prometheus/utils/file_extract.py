"""Shared file text extraction for Telegram gateway + Beacon web uploads.

Both gateways call extract_text() with a file path and optional MIME type.
Supported formats:
  - .txt, .md, .csv, .json, .yaml, .toml, .py, .js, .ts, .sh, .sql, etc. → UTF-8 read
  - .pdf → pypdfium2 (PDFium), plus the filled form fields its text layer leaves out
  - .docx → python-docx
  - .xlsx → openpyxl
"""

from __future__ import annotations

import ctypes
import logging
import threading
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_EXTRACT_BYTES = 30_000  # ~30 KB / ~7500 tokens — must fit alongside tools + system prompt

# Text-readable extensions (read as UTF-8)
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".toml",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".sh", ".sql",
    ".html", ".xml", ".log", ".ini", ".cfg", ".env",
    ".dockerfile", ".makefile", ".rs", ".go", ".java", ".c", ".cpp", ".h",
    ".swift", ".kt", ".rb", ".pl", ".lua", ".r", ".m",
}

# Binary formats with specialized extractors
BINARY_EXTRACTORS = {".pdf", ".docx", ".xlsx"}

SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | BINARY_EXTRACTORS

UNSUPPORTED_MSG = (
    "I can't read {ext} files yet. "
    "Supported formats: txt, md, csv, json, yaml, py, pdf, docx, xlsx "
    "(and other common text/code files)."
)


def extract_text(filepath: str, mime_type: str | None = None) -> str | None:
    """Extract text content from a file. Returns None if unsupported or empty.

    This is the single entry point — both Telegram and Beacon call this.
    """
    path = Path(filepath)
    if not path.is_file():
        logger.warning("extract_text: file not found: %s", filepath)
        return None

    ext = path.suffix.lower()

    # Size guard
    size = path.stat().st_size
    if size == 0:
        return None

    # Text-based files
    if ext in TEXT_EXTENSIONS:
        return _extract_plain_text(path)

    # PDF
    if ext == ".pdf":
        return _extract_pdf(path)

    # DOCX
    if ext == ".docx":
        return _extract_docx(path)

    # XLSX
    if ext == ".xlsx":
        return _extract_xlsx(path)

    return None


def is_supported(filepath: str) -> bool:
    """Check if a file extension is supported for text extraction."""
    return Path(filepath).suffix.lower() in SUPPORTED_EXTENSIONS


def unsupported_message(filepath: str) -> str:
    """Return a user-friendly message for unsupported file types."""
    ext = Path(filepath).suffix.lower()
    return UNSUPPORTED_MSG.format(ext=ext or "this type of")


def _extract_plain_text(path: Path) -> str | None:
    """Read a text file as UTF-8."""
    try:
        data = path.read_bytes()
        if len(data) > MAX_EXTRACT_BYTES:
            text = data[:MAX_EXTRACT_BYTES].decode("utf-8", errors="replace")
            return text + f"\n\n[... truncated at {MAX_EXTRACT_BYTES // 1000} KB]"
        return data.decode("utf-8", errors="replace")
    except Exception as exc:
        logger.warning("Failed to read text file %s: %s", path, exc)
        return None


# PDFium is not thread-safe (pypdfium2's documentation: "PDFium is inherently not thread-safe"), and uploads
# arrive concurrently from Telegram and Beacon. Every PDFium call, and the lifetime of every document, page
# and text page, sits inside this one lock.
_PDFIUM_LOCK = threading.Lock()


def _clean_pdfium_text(text: str) -> str:
    """PDFium joins a word it de-hyphenated at a line break and marks the join with U+0002 ("learn\x02ing"),
    and ends lines with CRLF. Dropping the marker gives the joined word; CRLF becomes LF."""
    return text.replace("\r\n", "\n").replace("\x02", "")


def _utf16(fn, *args) -> str:
    """Call a PDFium getter that fills a caller-sized UTF-16LE buffer and returns the byte count."""
    size = fn(*args, None, 0)
    if size <= 2:
        return ""
    buf = (ctypes.c_ushort * (size // 2))()
    fn(*args, buf, size)
    return bytes(buf).decode("utf-16-le", errors="replace").rstrip("\x00")


def _filled_fields(pdfium_raw, page, formenv) -> list[str]:
    """``name: value`` for each filled form field on the page. PDFium's text layer never holds what was typed
    into a form; without this a filled form loses exactly the part someone wrote. Empty fields and unchecked
    boxes ("Off") are left out as noise. The comparison (docs/audits/pdf-extraction/) found 8 of 8 values this
    way, against 6 of 8 from PyMuPDF's text."""
    fields: list[str] = []
    for index in range(pdfium_raw.FPDFPage_GetAnnotCount(page.raw)):
        annot = pdfium_raw.FPDFPage_GetAnnot(page.raw, index)
        if not annot:
            continue
        try:
            if pdfium_raw.FPDFAnnot_GetSubtype(annot) != pdfium_raw.FPDF_ANNOT_WIDGET:
                continue
            value = _utf16(pdfium_raw.FPDFAnnot_GetFormFieldValue, formenv, annot).strip()
            if not value or value == "Off":
                continue
            name = _utf16(pdfium_raw.FPDFAnnot_GetFormFieldName, formenv, annot).rsplit(".", 1)[-1]
            fields.append(f"{name}: {value}" if name else value)
        finally:
            pdfium_raw.FPDFPage_CloseAnnot(annot)
    return fields


def _extract_pdf(path: Path) -> str | None:
    """Extract text from a PDF with pypdfium2: each page's text, then that page's filled form fields."""
    try:
        import pypdfium2 as pdfium
        import pypdfium2.raw as pdfium_raw
    except ImportError:
        logger.warning("pypdfium2 is not installed — cannot extract PDF text (it is a core dependency)")
        return f"[PDF file: {path.name} — PDF text extraction is unavailable on this install]"

    try:
        with _PDFIUM_LOCK:
            pdf = pdfium.PdfDocument(str(path))
            try:
                formenv = None
                try:
                    pdf.init_forms()
                    formenv = getattr(pdf.formenv, "raw", None)
                except Exception:  # a PDF without a usable form is still a PDF
                    logger.debug("init_forms failed for %s", path.name, exc_info=True)
                page_count = len(pdf)
                pages: list[str] = []
                total_chars = 0
                for number in range(page_count):
                    page = pdf[number]
                    try:
                        textpage = page.get_textpage()
                        try:
                            text = _clean_pdfium_text(textpage.get_text_bounded())
                        finally:
                            textpage.close()
                        fields = _filled_fields(pdfium_raw, page, formenv) if formenv is not None else []
                    finally:
                        page.close()
                    if fields:
                        text += "\n[Form fields]\n" + "\n".join(fields)
                    if total_chars + len(text) > MAX_EXTRACT_BYTES:
                        pages.append(text[: MAX_EXTRACT_BYTES - total_chars])
                        pages.append(f"\n\n[... truncated at page {number + 1}/{page_count}]")
                        break
                    pages.append(text)
                    total_chars += len(text)
            finally:
                pdf.close()
        result = "\n".join(pages).strip()
        return result if result else f"[PDF file: {path.name} — no extractable text (may be image-based)]"
    except Exception as exc:
        logger.warning("Failed to extract PDF %s: %s", path, exc)
        return f"[PDF file: {path.name} — extraction failed: {exc}]"


def _extract_docx(path: Path) -> str | None:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx not installed — cannot extract DOCX text. Install: uv add python-docx")
        return f"[DOCX file: {path.name} — install python-docx to extract text]"

    try:
        doc = Document(str(path))
        paragraphs = []
        total_chars = 0
        for para in doc.paragraphs:
            text = para.text
            if total_chars + len(text) > MAX_EXTRACT_BYTES:
                paragraphs.append(f"\n[... truncated at {MAX_EXTRACT_BYTES // 1000} KB]")
                break
            paragraphs.append(text)
            total_chars += len(text)
        result = "\n".join(paragraphs).strip()
        return result if result else f"[DOCX file: {path.name} — no text content]"
    except Exception as exc:
        logger.warning("Failed to extract DOCX %s: %s", path, exc)
        return f"[DOCX file: {path.name} — extraction failed: {exc}]"


def _extract_xlsx(path: Path) -> str | None:
    """Extract text from XLSX using openpyxl."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("openpyxl not installed — cannot extract XLSX text. Install: uv add openpyxl")
        return f"[XLSX file: {path.name} — install openpyxl to extract text]"

    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
        sheets = []
        total_chars = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            rows.append(f"## Sheet: {sheet_name}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                line = " | ".join(cells)
                if total_chars + len(line) > MAX_EXTRACT_BYTES:
                    rows.append(f"[... truncated at {MAX_EXTRACT_BYTES // 1000} KB]")
                    total_chars = MAX_EXTRACT_BYTES + 1
                    break
                rows.append(line)
                total_chars += len(line) + 1
            sheets.append("\n".join(rows))
            if total_chars > MAX_EXTRACT_BYTES:
                break
        wb.close()
        result = "\n\n".join(sheets).strip()
        return result if result else f"[XLSX file: {path.name} — empty spreadsheet]"
    except Exception as exc:
        logger.warning("Failed to extract XLSX %s: %s", path, exc)
        return f"[XLSX file: {path.name} — extraction failed: {exc}]"
